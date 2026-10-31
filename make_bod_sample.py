"""Generate explicitly synthetic BOD data. Never mixes with real input folders."""
import argparse
import tempfile
from pathlib import Path
import pandas as pd
from docx import Document
from report import SCHEMAS


def make_sample(folder, month='2026-10'):
    folder=Path(folder)
    folder.mkdir(parents=True,exist_ok=True)
    if any(folder.iterdir()): raise ValueError('Thư mục mẫu phải trống để tránh trộn dữ liệu thật.')
    end=pd.Period(month).end_time.normalize()
    frames={k:[] for k in SCHEMAS}; layers=[]; movements=[]; plans=[]
    for period in pd.period_range(month[:4]+'-01',month,freq='M'):
        m=str(period)
        for zone in 'ABCDEF':
            entity='PHAP_NHAN_GIA_LAP_1' if zone in 'ABC' else 'PHAP_NHAN_GIA_LAP_2'
            movements.append([m,entity,zone,1000,0,0,0,0,1000])
            for date in pd.date_range(period.start_time,period.end_time.normalize()):
                frames['dan_ga'].append([date,entity,zone,1000])
                frames['trung'].append([date,entity,zone,800,760])
                frames['cam'].append([date,entity,zone,'CAM_GIA_LAP',120])
                layers.append([date,entity,zone,1000])
        days=period.days_in_month
        for metric,value,unit in [('tong_dan',6000,'con'),('tong_trung',days*6*850,'quả'),('trung_dat',days*6*810,'quả'),('cam_kg',days*6*118,'kg'),('ty_le_de_pct',85,'%'),('ty_le_dat_pct',96,'%')]:
            plans.append([m,metric,value,unit])
        for number in range(1,4):
            frames['kpi'].append([m,f'NV_GIA_{number:02d}',f'Nhân viên GIẢ LẬP {number}','Sản lượng giả lập',70,100,90+number*5,'tang'])
            frames['kpi'].append([m,f'NV_GIA_{number:02d}',f'Nhân viên GIẢ LẬP {number}','Tiêu hao giả lập',30,10,10,'giam'])
    with pd.ExcelWriter(folder/'SAN_XUAT_MAU_GIA_LAP.xlsx',engine='openpyxl') as writer:
        for key in ['dan_ga','trung','cam']:
            pd.DataFrame(frames[key],columns=SCHEMAS[key]).to_excel(writer,sheet_name=key,index=False)
        pd.DataFrame(layers,columns=['ngay','phap_nhan','khu','so_ga_de']).to_excel(writer,sheet_name='ga_de',index=False)
        pd.DataFrame(movements,columns=['thang','phap_nhan','khu','dau_ky','nhap','xuat','chet','loai','cuoi_ky']).to_excel(writer,sheet_name='bien_dong',index=False)
    with pd.ExcelWriter(folder/'KPI_KE_HOACH_MAU_GIA_LAP.xlsx',engine='openpyxl') as writer:
        pd.DataFrame(frames['kpi'],columns=SCHEMAS['kpi']).to_excel(writer,sheet_name='kpi',index=False)
        pd.DataFrame(plans,columns=['thang','chi_so','ke_hoach','don_vi']).to_excel(writer,sheet_name='ke_hoach',index=False)
    doc=Document();doc.add_heading('GHI CHÚ MẪU GIẢ LẬP — KHÔNG PHẢI DỮ LIỆU TAFA',0)
    table=doc.add_table(rows=1,cols=4)
    for c,t in zip(table.rows[0].cells,['thang','nhom','noi_dung','can_cu']):c.text=t
    rows=[('nguyen_nhan','Nguồn mẫu chưa có hồ sơ xác nhận nguyên nhân giảm sản lượng.','Mẫu giả lập; không khẳng định nguyên nhân'),
          ('giai_phap','Đối soát nhật ký gà đẻ và sản lượng theo từng khu.','Đề xuất minh họa từ chênh lệch kế hoạch trong mẫu'),
          ('kien_nghi','Phê duyệt việc đối soát bộ dữ liệu mẫu trước khi thử nghiệm chính thức.','Chỉ là kiến nghị thử nghiệm')]
    for row in rows:
        for c,t in zip(table.add_row().cells,(month,)+row):c.text=t
    doc.save(folder/'GHI_CHU_MAU_GIA_LAP.docx')

if __name__=='__main__':
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--directory',default=str(Path(tempfile.gettempdir())/'tafa-bod-demo'/'data'));p.add_argument('--month',default='2026-10')
    a=p.parse_args();make_sample(a.directory,a.month);print('Đã tạo 2 Excel và 1 Word hoàn toàn giả lập.')
