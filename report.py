"""TAFA monthly reporting. All input stays on the local filesystem."""
import argparse
from pathlib import Path
import math
import re
from local_runtime import discover_files, source_name, protect_official_input
import pandas as pd
from docx import Document
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, Reference

SCHEMAS = {
    'dan_ga': ['ngay','phap_nhan','khu','so_con'],
    'trung': ['ngay','phap_nhan','khu','tong_trung','trung_dat'],
    'cam': ['ngay','phap_nhan','khu','loai_cam','khoi_luong_kg'],
    'kpi': ['thang','ma_nv','ho_ten','chi_tieu','trong_so','muc_tieu','thuc_hien','chieu'],
}
NUMBERS = {'dan_ga':['so_con'], 'trung':['tong_trung','trung_dat'],
           'cam':['khoi_luong_kg'], 'kpi':['trong_so','muc_tieu','thuc_hien']}
KEYS = {'dan_ga':['ngay','phap_nhan','khu'],
        'trung':['ngay','phap_nhan','khu'],
        'cam':['ngay','phap_nhan','khu','loai_cam'],
        'kpi':['thang','ma_nv','chi_tieu']}


def read_inputs(folder, month, year_to_date=False, recursive=False, input_files=None):
    """Fail closed: no aggregation if any input error is found."""
    tables = {k: [] for k in SCHEMAS}
    issues = []
    def issue(level, source, message):
        issues.append({'muc_do':level,'nguon':source,'noi_dung':message})
    files = [p for p in (input_files if input_files is not None else discover_files(folder, recursive)) if p.suffix.lower() in {'.xlsx','.xlsm','.xls'}]
    if not files:
        issue('LOI',str(folder),'Không tìm thấy file Excel đầu vào.')
    for path in files:
        if path.suffix.lower() == '.xls':
            issue('LOI',path.name,'Định dạng .xls chưa hỗ trợ; lưu lại thành .xlsx.'); continue
        try:
            sheets = pd.read_excel(path, sheet_name=None, engine='openpyxl', dtype=object)
        except Exception as exc:
            issue('LOI',path.name,f'Không đọc được workbook ({type(exc).__name__}).'); continue
        for name, frame in sheets.items():
            key = name.strip().lower()
            source = f'{source_name(path, folder)}/{name}'
            if key not in SCHEMAS:
                issue('CANH_BAO',source,'Sheet ngoài mẫu, không đưa vào phân tích.'); continue
            frame.columns = [str(c).strip().lower() for c in frame.columns]
            missing = set(SCHEMAS[key]) - set(frame.columns)
            if missing or frame.columns.duplicated().any():
                issue('LOI',source,f'Thiếu cột {sorted(missing)} hoặc trùng tên cột.'); continue
            frame = frame.dropna(how='all').copy()
            if frame.empty:
                issue('LOI',source,'Sheet không có dữ liệu.'); continue
            frame['_nguon'] = source
            frame['_dong'] = frame.index + 2
            for col in SCHEMAS[key]:
                blank = frame[col].isna() | frame[col].astype(str).str.strip().eq('')
                for row in frame.loc[blank,'_dong']:
                    issue('LOI',f'{source}: dòng {row}',f'Thiếu giá trị {col}.')
            for col in NUMBERS[key]:
                frame[col] = pd.to_numeric(frame[col], errors='coerce')
                bad = frame[col].isna() | frame[col].map(lambda x: not math.isfinite(x) if pd.notna(x) else True) | (frame[col] < 0)
                if col in {'so_con','tong_trung','trung_dat'}:
                    bad |= frame[col].mod(1).ne(0)
                for row in frame.loc[bad,'_dong']:
                    issue('LOI',f'{source}: dòng {row}',f'{col} phải là số hữu hạn không âm; số con/trứng phải nguyên.')
            if key != 'kpi':
                # Require real Excel dates or ISO date strings to avoid day/month ambiguity.
                raw_dates = frame['ngay']
                ambiguous = raw_dates.map(lambda v: isinstance(v,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',v.strip()) is None)
                frame['ngay'] = pd.to_datetime(raw_dates, errors='coerce', format='mixed')
                frame.loc[ambiguous,'ngay'] = pd.NaT
                bad = frame['ngay'].isna() | frame['ngay'].ne(frame['ngay'].dt.normalize()) | (~frame['ngay'].dt.strftime('%Y-%m').between(month[:4] + '-01', month) if year_to_date else frame['ngay'].dt.strftime('%Y-%m').ne(month))
                for row in frame.loc[bad,'_dong']:
                    issue('LOI',f'{source}: dòng {row}','Ngày không hợp lệ hoặc ngoài tháng báo cáo.')
                for col in ['phap_nhan','khu']:
                    frame[col] = frame[col].astype(str).str.strip()
                frame['khu'] = frame['khu'].str.upper()
                if not frame['khu'].isin(list('ABCDEF')).all():
                    issue('LOI',source,'Khu phải thuộc A–F.')
            else:
                for col in ['thang','ma_nv','ho_ten','chi_tieu','chieu']:
                    frame[col] = frame[col].astype(str).str.strip()
                if not frame['thang'].str.fullmatch(r'\d{4}-\d{2}').all() or not (frame['thang'].between(month[:4] + '-01', month) if year_to_date else frame['thang'].eq(month)).all():
                    issue('LOI',source,'thang phải có dạng YYYY-MM và đúng tháng báo cáo.')
                if not frame['chieu'].isin(['tang','giam']).all():
                    issue('LOI',source,'chieu phải là tang hoặc giam.')
                if (frame['muc_tieu'] <= 0).any() or ((frame['chieu'] == 'giam') & (frame['thuc_hien'] <= 0)).any():
                    issue('LOI',source,'Mục tiêu phải > 0; KPI chiều giảm yêu cầu thực hiện > 0.')
            if key == 'trung' and (frame['trung_dat'] > frame['tong_trung']).any():
                issue('LOI',source,'Trứng đạt không được lớn hơn tổng trứng.')
            tables[key].append(frame)
    result = {}
    for key, parts in tables.items():
        if not parts:
            issue('LOI',key,'Thiếu nhóm dữ liệu bắt buộc.'); continue
        df = pd.concat(parts, ignore_index=True)
        duplicates = df.duplicated(KEYS[key], keep=False)
        for _, row in df.loc[duplicates].iterrows():
            issue('LOI',f"{row['_nguon']}: dòng {row['_dong']}",'Trùng khóa dữ liệu, kể cả giữa các file; cần đối soát.')
        result[key] = df
    if 'kpi' in result:
        df = result['kpi']
        for employee, part in df.groupby(['thang','ma_nv']):
            if part['ho_ten'].nunique() != 1:
                issue('LOI','kpi',f'Mã {employee} có nhiều tên nhân viên.')
            if not math.isclose(part['trong_so'].sum(),100,abs_tol=0.001):
                issue('LOI','kpi',f'Tổng trọng số của {employee} phải bằng 100%.')
    return result, pd.DataFrame(issues, columns=['muc_do','nguon','noi_dung'])


def analyze(data, month):
    herd = data['dan_ga']
    # Herd is a stock, never sum successive daily snapshots.
    latest = herd['ngay'].max()
    end = herd[herd['ngay'].eq(latest)]
    herd_summary = end.groupby(['phap_nhan','khu'],as_index=False)['so_con'].sum()
    eggs = data['trung'].groupby(['phap_nhan','khu'],as_index=False)[['tong_trung','trung_dat']].sum()
    eggs['trung_khong_dat'] = eggs['tong_trung'] - eggs['trung_dat']
    eggs['ty_le_dat_pct'] = eggs['trung_dat'].div(eggs['tong_trung'].replace(0,float('nan'))) * 100
    feed = data['cam'].groupby(['phap_nhan','khu','loai_cam'],as_index=False)['khoi_luong_kg'].sum()
    kpi = data['kpi'].copy()
    kpi['hoan_thanh_pct'] = kpi['thuc_hien'] / kpi['muc_tieu'] * 100
    decreasing = kpi['chieu'].eq('giam')
    kpi.loc[decreasing,'hoan_thanh_pct'] = kpi.loc[decreasing,'muc_tieu'] / kpi.loc[decreasing,'thuc_hien'] * 100
    kpi['diem_trong_so'] = kpi['hoan_thanh_pct'] * kpi['trong_so'] / 100
    ranking = kpi.groupby(['ma_nv','ho_ten'],as_index=False)['diem_trong_so'].sum()
    ranking['hang'] = ranking['diem_trong_so'].rank(method='min',ascending=False).astype(int)
    ranking = ranking.sort_values(['hang','ma_nv'])
    notes = [f'Tháng báo cáo: {month}. Chỉ tổng hợp dữ liệu được cung cấp; không xác nhận độ đầy đủ.',
             f'Tổng đàn là ảnh chụp chung ngày {latest:%Y-%m-%d}; không cộng đàn qua các ngày.',
             'Trứng và cám là phát sinh theo ngày; không dùng số lũy kế hoặc dòng tổng cộng.',
             'Chất lượng trứng: trứng đạt / tổng trứng; tổng bằng 0 thì tỷ lệ để trống.',
             'KPI: chiều tăng = thực hiện/mục tiêu; chiều giảm = mục tiêu/thực hiện. Không giới hạn điểm trần.',
             'Điểm tổng = tổng tỷ lệ hoàn thành × trọng số/100; đồng điểm đồng hạng (1, 1, 3).',
             'Chính sách KPI cần được phê duyệt trước khi sử dụng cho quyết định nhân sự.']
    expected = set(map(tuple,herd[['phap_nhan','khu']].drop_duplicates().to_numpy()))
    present = set(map(tuple,end[['phap_nhan','khu']].to_numpy()))
    if expected != present:
        notes.append('CẢNH BÁO: ngày chốt thiếu pháp nhân/khu đã xuất hiện trong các ngày trước; tổng đàn chưa đủ phạm vi.')
    if latest != pd.Period(month).end_time.normalize():
        notes.append('CẢNH BÁO: chưa có ảnh chụp đàn ngày cuối tháng.')
    return {'Tong_dan':herd_summary,'San_luong_trung':eggs,'Su_dung_cam':feed,
            'Xep_hang_KPI':ranking,'Chi_tiet_KPI':kpi.drop(columns=['_nguon','_dong']),
            'Ghi_chu':pd.DataFrame({'noi_dung':notes})}


def export_report(tables, issues, output, month, sample=False):
    output.mkdir(parents=True,exist_ok=True)
    if sample:
        tables['Ghi_chu'].loc[len(tables['Ghi_chu'])] = 'DỮ LIỆU MẪU GIẢ LẬP — KHÔNG SỬ DỤNG RA QUYẾT ĐỊNH'
    suffix = '_MAU_GIA_LAP' if sample else ''
    stem = output / f'TAFA_{month}{suffix}'
    if stem.with_suffix('.xlsx').exists() or stem.with_suffix('.docx').exists():
        raise FileExistsError('Báo cáo đã tồn tại; chọn thư mục output mới để bảo toàn kết quả cũ.')
    with pd.ExcelWriter(stem.with_suffix('.xlsx'),engine='openpyxl') as writer:
        for name, df in {**tables,'Kiem_tra_du_lieu':issues}.items():
            df.to_excel(writer,sheet_name=name,index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill('solid',fgColor='17365D'); cell.font = Font(color='FFFFFF',bold=True)
            for column in ws.columns:
                ws.column_dimensions[column[0].column_letter].width = min(75,max(16,max(len(str(c.value or '')) for c in column)+2))
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical='top',wrap_text=True)
                    if isinstance(cell.value,float): cell.number_format = '#,##0.00'
                    # Prevent Excel formula injection from source text.
                    if isinstance(cell.value,str): cell.data_type = 's'
        ws = writer.sheets['Tong_dan']
        chart = BarChart(); chart.title = 'Đàn gà theo pháp nhân / khu'
        chart.add_data(Reference(ws,min_col=3,max_col=3,min_row=1,max_row=ws.max_row),titles_from_data=True)
        chart.set_categories(Reference(ws,min_col=1,max_col=2,min_row=2,max_row=ws.max_row))
        ws.add_chart(chart,'F2')
    doc = Document()
    section = doc.sections[0]; section.header.paragraphs[0].text = 'TAFA | PHÒNG CHĂN NUÔI | LƯU HÀNH NỘI BỘ'
    doc.add_heading('BÁO CÁO PHÒNG CHĂN NUÔI',0)
    doc.add_paragraph(f'Tháng {month} • Kính trình Chủ tịch Hội đồng quản trị')
    if sample: doc.add_paragraph('DỮ LIỆU MẪU GIẢ LẬP — KHÔNG SỬ DỤNG RA QUYẾT ĐỊNH',style='Subtitle')
    doc.add_heading('Tóm tắt điều hành',1)
    eggs = tables['San_luong_trung']; total = eggs['tong_trung'].sum(); good = eggs['trung_dat'].sum()
    rate = f'{good/total*100:,.2f}%' if total else 'Không xác định (tổng trứng = 0)'
    for line in [f"Đàn gà tại ngày chốt: {tables['Tong_dan']['so_con'].sum():,.0f} con.",
                 f'Tổng trứng: {total:,.0f}; trứng đạt: {good:,.0f}; tỷ lệ đạt: {rate}.',
                 f"Cám sử dụng: {tables['Su_dung_cam']['khoi_luong_kg'].sum():,.2f} kg.",
                 f"Số nhân viên có KPI: {len(tables['Xep_hang_KPI'])}."]:
        doc.add_paragraph(line)
    titles = {'Tong_dan':'1. Tổng đàn gà','San_luong_trung':'2. Sản lượng và chất lượng trứng',
              'Su_dung_cam':'3. Sử dụng cám','Xep_hang_KPI':'4. KPI và xếp hạng nhân viên'}
    for key,title in titles.items():
        doc.add_heading(title,1)
        df = tables[key]; table = doc.add_table(rows=1,cols=len(df.columns)); table.style = 'Light Shading Accent 1'
        for c,col in zip(table.rows[0].cells,df.columns): c.text = col
        for row in df.itertuples(index=False,name=None):
            for cell,value in zip(table.add_row().cells,row):
                cell.text = '' if pd.isna(value) else (f'{value:,.2f}' if isinstance(value,float) else str(value))
    doc.add_heading('Phạm vi, phương pháp và nội dung cần đối soát',1)
    for note in tables['Ghi_chu']['noi_dung']: doc.add_paragraph(note)
    for row in issues.itertuples(index=False): doc.add_paragraph(f'{row.muc_do}: {row.nguon}: {row.noi_dung}')
    section.footer.paragraphs[0].text = 'Tài liệu nội bộ • Kiểm tra và phê duyệt trước khi phát hành'
    doc.save(stem.with_suffix('.docx'))
    return stem


def run(folder, output, month='2026-10', sample=False, recursive=False, confirmed_real=False):
    if re.fullmatch(r'\d{4}-\d{2}',month) is None:
        raise ValueError('Tháng phải có dạng YYYY-MM.')
    pd.Period(month,freq='M')
    output = Path(output); output.mkdir(parents=True,exist_ok=True)
    files = discover_files(folder, recursive)
    data, issues = read_inputs(folder, month, recursive=recursive, input_files=files)
    if issues['muc_do'].eq('LOI').any():
        # Diagnostic file contains source names and employee codes: also keep private.
        issues.to_csv(output/'loi_du_lieu.csv',index=False,encoding='utf-8-sig')
        return False
    if not sample:
        protect_official_input(files, confirmed_real)
    export_report(analyze(data,month),issues,output,month,sample)
    return True


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data',default='data'); parser.add_argument('--output',default='output')
    parser.add_argument('--month',default='2026-10'); parser.add_argument('--sample',action='store_true')
    parser.add_argument('--recursive',action='store_true'); parser.add_argument('--confirm-real-data',action='store_true')
    args = parser.parse_args()
    try:
        if not run(args.data,args.output,args.month,args.sample,args.recursive,args.confirm_real_data):
            print('Dữ liệu chưa hợp lệ. Xem loi_du_lieu.csv trong thư mục output. Chưa tạo báo cáo.'); return 2
    except (ValueError,OSError) as exc:
        print(f'Không thể tạo báo cáo: {exc}'); return 2
    print('Đã tạo báo cáo Excel và Word trong thư mục output đã chọn.'); return 0

if __name__ == '__main__':
    raise SystemExit(main())
