"""Local, evidence-based BOD reporting. No network calls or inferred observations."""
import argparse
import hashlib
import json
import math
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
import pandas as pd
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.chart.data import CategoryChartData
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from report import read_inputs, analyze
from local_runtime import discover_files, source_name, protect_official_input, select_font, find_office, read_config

GREEN='145C43'; GOLD='C6A34B'; DARK='19392F'; MUTED='596B63'; LIGHT='F4F7F3'
METRICS = {
 'tong_dan':('Đàn tại ngày chốt','con','tang'),
 'tong_trung':('Sản lượng trứng','quả','tang'),
 'trung_dat':('Trứng đạt','quả','tang'),
 'cam_kg':('Cám sử dụng','kg','giam'),
 'ty_le_de_pct':('Tỷ lệ đẻ','%','tang'),
 'ty_le_dat_pct':('Tỷ lệ trứng đạt','%','tang'),
}
EXTRA = {
 'ga_de':['ngay','phap_nhan','khu','so_ga_de'],
 'bien_dong':['thang','phap_nhan','khu','dau_ky','nhap','xuat','chet','loai','cuoi_ky'],
 'ke_hoach':['thang','chi_so','ke_hoach','don_vi'],
}
NARRATIVE = {'ton_tai','nguyen_nhan','giai_phap','kien_nghi'}


def fmt(value, unit=''):
    if value is None or pd.isna(value): return 'Chưa đủ căn cứ'
    return f'{value:,.2f}' + (f' {unit}' if unit else '')


def load(folder, month, recursive=False, input_files=None):
    if re.fullmatch(r'\d{4}-\d{2}',month) is None:
        raise ValueError('Tháng phải là YYYY-MM.')
    pd.Period(month,freq='M')
    folder=Path(folder)
    if not folder.is_dir(): raise ValueError('Không tìm thấy thư mục data.')
    files=input_files if input_files is not None else discover_files(folder,recursive)
    base, issues=read_inputs(folder,month,year_to_date=True,recursive=recursive,input_files=files)
    # These sheets belong to the BOD schema and are validated below.
    mask=issues.noi_dung.eq('Sheet ngoài mẫu, không đưa vào phân tích.') & issues.nguon.map(lambda x: x.rsplit('/',1)[-1].strip().lower() in EXTRA)
    issues=issues.loc[~mask].to_dict('records')
    extra={k:[] for k in EXTRA}; narratives=[]; sources=[]
    def error(source,message): issues.append(dict(muc_do='LOI',nguon=source,noi_dung=message))
    for path in files:
        relative=source_name(path, folder)
        if path.name.startswith('~$') or path.suffix.lower() not in {'.xlsx','.xlsm','.docx'}: continue
        sources.append({'file':relative,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
        if path.suffix.lower()=='.docx':
            try:
                doc=Document(path)
                for paragraph in doc.paragraphs:
                    if paragraph.text.strip():
                        narratives.append({'nhom':'tham_khao','noi_dung':paragraph.text.strip(),'can_cu':'Văn bản nguồn; chưa xác minh','nguon':relative})
                for table in doc.tables:
                    headers=[c.text.strip().lower() for c in table.rows[0].cells]
                    if not {'thang','nhom','noi_dung','can_cu'}.issubset(headers):
                        narratives.append({'nhom':'tham_khao','noi_dung':'\n'.join(' | '.join(c.text for c in r.cells) for r in table.rows),'can_cu':'Bảng chưa đúng mẫu; không trích số liệu','nguon':relative}); continue
                    for row in table.rows[1:]:
                        item=dict(zip(headers,[c.text.strip() for c in row.cells])); item['nguon']=relative
                        if item.get('nhom') not in NARRATIVE or not item.get('noi_dung') or not item.get('can_cu'):
                            error(relative,'Bảng Word cần nhom hợp lệ, noi_dung và can_cu.'); continue
                        if item.get('thang')!=month:
                            item['nhom_goc']=item['nhom'];item['nhom']='tham_khao'
                        narratives.append(item)
            except Exception as exc: error(relative,f'Không đọc được Word ({type(exc).__name__}).')
            continue
        try: sheets=pd.read_excel(path,sheet_name=None,engine='openpyxl',dtype=object)
        except Exception: continue # Base reader already reports the workbook error.
        for name,df in sheets.items():
            key=name.strip().lower()
            if key not in EXTRA: continue
            source=f'{relative}/{name}'
            df.columns=[str(c).strip().lower() for c in df.columns]
            if not set(EXTRA[key]).issubset(df.columns) or df.columns.duplicated().any():
                error(source,f'Thiếu hoặc trùng cột: yêu cầu {EXTRA[key]}.'); continue
            df=df.dropna(how='all').copy()
            if df.empty: error(source,'Sheet bổ sung trống.'); continue
            for col in EXTRA[key]:
                if (df[col].isna() | df[col].astype(str).str.strip().eq('')).any(): error(source,f'Thiếu giá trị {col}.')
            nums={'ga_de':['so_ga_de'],'bien_dong':['dau_ky','nhap','xuat','chet','loai','cuoi_ky'],'ke_hoach':['ke_hoach']}[key]
            for col in nums:
                df[col]=pd.to_numeric(df[col],errors='coerce')
                bad=df[col].isna() | ~df[col].map(lambda x: math.isfinite(x) if pd.notna(x) else False) | df[col].lt(0)
                if key!='ke_hoach': bad |= df[col].mod(1).ne(0)
                if bad.any(): error(source,f'{col} phải là số hữu hạn không âm; số con phải nguyên.')
            if key=='ga_de':
                raw=df.ngay
                invalid=raw.map(lambda x: isinstance(x,str) and re.fullmatch(r'\d{4}-\d{2}-\d{2}',x.strip()) is None)
                df['ngay']=pd.to_datetime(raw,format='mixed',errors='coerce'); df.loc[invalid,'ngay']=pd.NaT
                periods=df.ngay.dt.strftime('%Y-%m')
                if (df.ngay.isna() | df.ngay.ne(df.ngay.dt.normalize()) | ~periods.between(month[:4]+'-01',month)).any(): error(source,'Ngày gà đẻ không hợp lệ hoặc ngoài kỳ.')
            else:
                df['thang']=df.thang.astype(str).str.strip()
                if not df.thang.str.fullmatch(r'\d{4}-\d{2}').all() or not df.thang.between(month[:4]+'-01',month).all(): error(source,'Tháng bổ sung không hợp lệ hoặc ngoài kỳ.')
            if key in {'ga_de','bien_dong'}:
                df['khu']=df.khu.astype(str).str.strip().str.upper()
                df['phap_nhan']=df.phap_nhan.astype(str).str.strip()
                if not df.khu.isin(list('ABCDEF')).all(): error(source,'Khu phải thuộc A–F.')
            if key=='ke_hoach':
                for col in ['chi_so','don_vi']: df[col]=df[col].astype(str).str.strip()
                for _,r in df.iterrows():
                    if r.chi_so not in METRICS: error(source,f'Chỉ số kế hoạch không hỗ trợ: {r.chi_so}.')
                    elif r.don_vi!=METRICS[r.chi_so][1]: error(source,f'Sai đơn vị của {r.chi_so}.')
                    if r.chi_so in {'ty_le_de_pct','ty_le_dat_pct'} and r.ke_hoach>100: error(source,'Kế hoạch tỷ lệ không vượt 100%.')
            if key=='bien_dong':
                if not (df.dau_ky+df.nhap-df.xuat-df.chet-df.loai).eq(df.cuoi_ky).all(): error(source,'Không cân đối: đầu kỳ + nhập − xuất − chết − loại ≠ cuối kỳ.')
            df['_nguon']=source; extra[key].append(df)
    for key,parts in extra.items():
        extra[key]=pd.concat(parts,ignore_index=True) if parts else pd.DataFrame(columns=EXTRA[key]+['_nguon'])
        keys={'ga_de':['ngay','phap_nhan','khu'],'bien_dong':['thang','phap_nhan','khu'],'ke_hoach':['thang','chi_so']}[key]
        if extra[key].duplicated(keys).any(): error(key,'Trùng khóa giữa các dòng/file.')
    # Required base groups must also have data in the selected report month.
    for key,df in base.items():
        selected=df.thang.eq(month) if key=='kpi' else df.ngay.dt.strftime('%Y-%m').eq(month)
        if not selected.any(): error(key,'Không có dữ liệu của tháng báo cáo.')
    return base,extra,narratives,pd.DataFrame(issues,columns=['muc_do','nguon','noi_dung']),sources


def subset(base,month):
    return {k:df.loc[df.thang.eq(month) if k=='kpi' else df.ngay.dt.strftime('%Y-%m').eq(month)].copy() for k,df in base.items()}


def laying_rate(eggs,layers):
    if eggs.empty or layers.empty: return None
    keys=['ngay','phap_nhan','khu']
    joined=eggs[keys+['tong_trung']].merge(layers[keys+['so_ga_de']],on=keys,how='outer',indicator=True)
    if not joined['_merge'].eq('both').all() or joined.so_ga_de.sum()<=0: return None
    if ((joined.so_ga_de==0) & (joined.tong_trung>0)).any(): return None
    return joined.tong_trung.sum()/joined.so_ga_de.sum()*100


def build_analysis(base,extra,month):
    current=subset(base,month); summary=analyze(current,month)
    eggs=current['trung']; feed=current['cam']; layers=extra['ga_de']
    layers=layers.loc[layers.ngay.dt.strftime('%Y-%m').eq(month)] if not layers.empty else layers
    total=eggs.tong_trung.sum(); good=eggs.trung_dat.sum()
    metrics={'tong_dan':summary['Tong_dan'].so_con.sum(),'tong_trung':total,'trung_dat':good,
             'cam_kg':feed.khoi_luong_kg.sum(),'ty_le_de_pct':laying_rate(eggs,layers),
             'ty_le_dat_pct':good/total*100 if total else None}
    comparisons=[]
    plans=extra['ke_hoach']; plans=plans.loc[plans.thang.eq(month)]
    for key,(label,unit,direction) in METRICS.items():
        rows=plans.loc[plans.chi_so.eq(key)]
        plan=rows.iloc[0].ke_hoach if not rows.empty else None; actual=metrics[key]
        delta=actual-plan if actual is not None and plan is not None else None
        comparisons.append({'chi_so':key,'ten':label,'don_vi':unit,'thuc_te':actual,'ke_hoach':plan,'chenh_lech':delta,'don_vi_chenh_lech':('điểm %' if unit=='%' else unit),
                            'hoan_thanh_pct':actual/plan*100 if actual is not None and plan is not None and plan>0 else None,
                            'danh_gia':('Chưa đủ căn cứ' if delta is None else ('Đạt' if (delta>=0 if direction=='tang' else delta<=0) else 'Cần đối soát'))})
    coverage=[]
    for key in ['dan_ga','trung','cam']:
        for (entity,zone),df in current[key].groupby(['phap_nhan','khu']):
            days=df.ngay.dt.normalize().nunique(); expected=pd.Period(month).days_in_month
            coverage.append({'nhom':key,'phap_nhan':entity,'khu':zone,'ngay_co_du_lieu':days,'ngay_trong_thang':expected})
    ytd=[]
    for period in pd.period_range(month[:4]+'-01',month,freq='M'):
        m=str(period); data=subset(base,m); e=data['trung']; f=data['cam']
        monthly_plans=extra['ke_hoach'].loc[extra['ke_hoach'].thang.eq(m)].set_index('chi_so')
        ytd.append({'ke_hoach_trung': monthly_plans.loc['tong_trung','ke_hoach'] if 'tong_trung' in monthly_plans.index else None,
                    'ke_hoach_cam_kg': monthly_plans.loc['cam_kg','ke_hoach'] if 'cam_kg' in monthly_plans.index else None,'thang':m,'trung_qua':e.tong_trung.sum() if not e.empty else None,
                    'trung_dat_qua':e.trung_dat.sum() if not e.empty else None,
                    'cam_kg':f.khoi_luong_kg.sum() if not f.empty else None,
                    'ngay_trung':e.ngay.dt.normalize().nunique(),'ngay_cam':f.ngay.dt.normalize().nunique(),
                    'ngay_thang':period.days_in_month})
    movements=extra['bien_dong'].loc[extra['bien_dong'].thang.eq(month)]
    observations=[]; solutions=[]
    for row in comparisons:
        if row['danh_gia']=='Cần đối soát':
            observations.append(f"{row['ten']}: chênh {fmt(row['chenh_lech'],row['don_vi_chenh_lech'])} so với kế hoạch. Chưa đủ căn cứ xác định nguyên nhân.")
            solutions.append(f"Đề xuất: đối soát {row['ten'].lower()} theo ngày/khu và xác minh nguyên nhân trước khi giao hành động.")
    if not coverage or any(r['ngay_co_du_lieu']<r['ngay_trong_thang'] for r in coverage):
        observations.append('Dữ liệu tháng chưa phủ đủ ngày ở một số nhóm/khu; tổng phản ánh phần nguồn đã cung cấp.')
        solutions.append('Đề xuất: hoàn thiện dữ liệu ngày và xác nhận phạm vi pháp nhân/khu trước khi phát hành.')
    if metrics['ty_le_de_pct'] is None:
        observations.append('Chưa tính được tỷ lệ đẻ: thiếu số gà đẻ-ngày hoặc khóa ngày/pháp nhân/khu không khớp.')
        solutions.append('Đề xuất: bổ sung số gà đẻ thực tế theo ngày, cùng phạm vi số trứng.')
    if metrics['ty_le_de_pct'] is not None and metrics['ty_le_de_pct']>100:
        observations.append('Tỷ lệ đẻ vượt 100%; cần kiểm tra mẫu số, phạm vi và đơn vị trứng.')
    if movements.empty:
        observations.append('Thiếu bảng biến động: chưa đủ căn cứ giải thích tăng/giảm tổng đàn.')
    else:
        closing=current['dan_ga'].loc[current['dan_ga'].ngay.eq(pd.Period(month).end_time.normalize())]
        joined=movements.merge(closing,on=['phap_nhan','khu'],how='outer',indicator=True)
        if not joined['_merge'].eq('both').all() or not joined.cuoi_ky.eq(joined.so_con).all():
            observations.append('Cuối kỳ bảng biến động chưa khớp đầy đủ ảnh chụp đàn cuối tháng; cần đối soát.')
    observations.extend(summary['Ghi_chu'].noi_dung[summary['Ghi_chu'].noi_dung.str.startswith('CẢNH BÁO')].tolist())
    return {'summary':summary,'metrics':metrics,'comparison':pd.DataFrame(comparisons),
            'coverage':pd.DataFrame(coverage),'ytd':pd.DataFrame(ytd),'movements':movements,
            'observations':observations,'solutions':solutions,'current':current}


class Deck:
    def __init__(self,month,sample,font=None):
        self.prs=Presentation(); self.prs.slide_width=Inches(13.333); self.prs.slide_height=Inches(7.5)
        self.month=month; self.sample=sample; self.checks=[]
        self.font=font or select_font()

    def box(self,slide,x,y,w,h,text,size=20,color=DARK,bold=False,fill=None):
        shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        if fill:
            shape.fill.solid(); shape.fill.fore_color.rgb=RGBColor.from_string(fill)
        tf=shape.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=Inches(.10); tf.margin_top=tf.margin_bottom=Inches(.015)
        # Conservative geometric guard. Full narrative is retained in the audit workbook.
        limit=max(12,int((w*72-15)/(size*.62)))
        import textwrap
        lines=sum(max(1,len(textwrap.wrap(line,width=limit,break_long_words=True))) for line in str(text).split('\n'))
        while lines*size*1.25>h*72-3 and size>12:
            size-=1; limit=max(12,int((w*72-15)/(size*.62)))
            lines=sum(max(1,len(textwrap.wrap(line,width=limit))) for line in str(text).split('\n'))
        if lines*size*1.25>h*72-3: raise ValueError(f'Nội dung vượt khung: {str(text)[:60]}; xem phụ lục.')
        for i,line in enumerate(str(text).split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line
            p.font.name=self.font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color)
        return shape

    def slide(self,title,message,source='Nguồn: Excel đầu vào; xem Doi_soat_BOD.xlsx để truy vết.'):
        slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb=RGBColor.from_string(LIGHT)
        accent=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,self.prs.slide_width,Inches(.045))
        accent.fill.solid();accent.fill.fore_color.rgb=RGBColor.from_string(GOLD);accent.line.fill.background()
        self.box(slide,.45,.15,12.4,.35,'TAFA GROUP  /  PHÒNG CHĂN NUÔI',12,GREEN,True)
        self.box(slide,.45,.63,12.4,.65,title,28,GREEN,True)
        self.box(slide,.45,1.40,12.4,.7,message,19,DARK,True)
        self.box(slide,.45,6.65,11.8,.43,source,10,MUTED)
        label='MẪU GIẢ LẬP • ' if self.sample else 'NỘI BỘ • '
        self.box(slide,.45,7.08,11.8,.24,f'{label}{self.month} | Phê duyệt trước khi phát hành',9,MUTED)
        self.box(slide,12.35,7.04,.5,.3,str(len(self.prs.slides)),10,GOLD,True)
        return slide

    def cards(self,slide,cards):
        count=len(cards); width=12.35/count
        for i,(label,value,detail) in enumerate(cards):
            x=.48+i*width
            self.box(slide,x,2.5,width-.18,.55,label,16,GREEN,True)
            self.box(slide,x,3.1,width-.18,1.1,value,27,GREEN,True,fill='FFFFFF')
            self.box(slide,x,4.45,width-.18,1.3,detail,15,MUTED)

    def bullets(self,slide,items):
        items=items or ['Chưa đủ căn cứ từ dữ liệu đầu vào.']
        # Exactly four readable cards maximum; all source text is in the audit file.
        shown=items[:4]
        for i,item in enumerate(shown):
            if len(item)>240: item=item[:220]+'… [xem phụ lục]'
            self.box(slide,.65,2.30+i*1.02,12,.91,f'{i+1:02d}  {item}',17,DARK,fill='FFFFFF')
        if len(items)>4:
            self.box(slide,.65,6.40,12,.23,f'Hiển thị 4/{len(items)} nội dung; đầy đủ trong Doi_soat_BOD.xlsx.',10,MUTED)

    def chart(self,slide,categories,series,unit,kind=XL_CHART_TYPE.COLUMN_CLUSTERED):
        if not categories:
            self.bullets(slide,['Chưa đủ căn cứ để lập biểu đồ.']); return
        chart_data=CategoryChartData()
        labels=[str(x).replace(' / ','\n') for x in categories]
        # XlsxWriter otherwise interprets category strings beginning '=' as formulas.
        chart_data.categories=[('\u200b'+label if label.lstrip().startswith(('=','+','-','@')) else label) for label in labels]
        for label,values in series: chart_data.add_series(label,[None if pd.isna(v) else float(v) for v in values])
        chart=slide.shapes.add_chart(kind,Inches(.65),Inches(2.35),Inches(12),Inches(4.05),chart_data).chart
        chart.has_legend=len(series)>1
        if chart.has_legend: chart.legend.position=XL_LEGEND_POSITION.BOTTOM
        chart.value_axis.has_title=True; chart.value_axis.axis_title.text_frame.text=unit
        chart.value_axis.major_gridlines.format.line.color.rgb=RGBColor.from_string('DCE5DE')
        chart.value_axis.major_gridlines.format.line.width=Pt(.5)
        chart.value_axis.minimum_scale=0
        chart.value_axis.tick_labels.font.size=Pt(11); chart.category_axis.tick_labels.font.size=Pt(10)
        for axis in [chart.value_axis,chart.category_axis]: axis.tick_labels.font.name=self.font
        for i,serie in enumerate(chart.series):
            serie.format.fill.solid(); serie.format.fill.fore_color.rgb=RGBColor.from_string([GREEN,GOLD,'7CA68B'][i%3])
            serie.format.line.color.rgb=RGBColor.from_string([GREEN,GOLD,'7CA68B'][i%3])
        if len(categories)<=6:
            chart.plots[0].has_data_labels=True; chart.plots[0].data_labels.font.size=Pt(11)

    def table(self,slide,headers,rows):
        rows=rows[:8]
        table=slide.shapes.add_table(len(rows)+1,len(headers),Inches(.65),Inches(2.35),Inches(12),Inches(3.95)).table
        for i,row in enumerate([headers]+rows):
            for j,value in enumerate(row):
                cell=table.cell(i,j); text=str(value)
                if len(text)>52: text=text[:35]+'… [phụ lục]'
                cell.text=text; cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                cell.fill.solid(); cell.fill.fore_color.rgb=RGBColor.from_string(GREEN if i==0 else ('FFFFFF' if i%2 else 'E9F0E9'))
                for p in cell.text_frame.paragraphs:
                    p.font.name=self.font;p.font.size=Pt(13 if i else 12);p.font.bold=i==0;p.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else DARK)

    def validate(self):
        if len(self.prs.slides)!=14: raise ValueError('PowerPoint phải có 14 slide.')
        for i,slide in enumerate(self.prs.slides,1):
            for shape in slide.shapes:
                if shape.left<0 or shape.top<0 or shape.left+shape.width>self.prs.slide_width or shape.top+shape.height>self.prs.slide_height:
                    raise ValueError(f'Slide {i}: thành phần ra ngoài khung.')
        self.checks=['14 slide','Các thành phần nằm trong khung','Khung chữ ước lượng theo cỡ font; chưa thay thế kiểm tra render trong PowerPoint']


def grouped_chart(df,keys,value,limit=8):
    grouped=df.groupby(keys)[value].sum().sort_values(ascending=False)
    labels=[' / '.join(map(str,k)) if isinstance(k,tuple) else str(k) for k in grouped.index]
    values=grouped.tolist()
    if len(values)>limit:
        return labels[:limit-1]+['Các nhóm khác'],values[:limit-1]+[sum(values[limit-1:])]
    return labels,values


def create_deck(a,extra,narratives,month,sample,font=None):
    deck=Deck(month,sample,font); m=a['metrics']; summary=a['summary']; current=a['current']
    slide=deck.slide('BÁO CÁO BOD | PHÒNG CHĂN NUÔI','Kết quả tháng và lũy kế được tổng hợp từ nguồn có thể truy vết.')
    if sample: deck.box(slide,.75,5.8,11.7,.5,'DỮ LIỆU MẪU GIẢ LẬP — KHÔNG SỬ DỤNG RA QUYẾT ĐỊNH',17,GREEN,True)
    deck.box(slide,.75,2.6,11.7,1.15,f'THÁNG {month[5:]}/{month[:4]}',42,GREEN,True)
    deck.box(slide,.75,4.1,11.7,1.6,'Kính trình Ban Lãnh đạo TAFA Group\nPhạm vi: nguồn đầu vào đã cung cấp • Không suy diễn số liệu còn thiếu',23,DARK)
    egg_comparison=a['comparison'].set_index('chi_so').loc['tong_trung']
    headline=(f"Trứng đạt {fmt(egg_comparison.hoan_thanh_pct,'%')} kế hoạch; chênh {fmt(egg_comparison.chenh_lech,'quả')} theo nguồn."
              if pd.notna(egg_comparison.hoan_thanh_pct) else 'Chưa có kế hoạch đủ căn cứ để đánh giá mức độ hoàn thành sản lượng.')
    slide=deck.slide('01 | Executive Summary',headline)
    deck.cards(slide,[('Đàn ngày chốt',fmt(m['tong_dan'],'con'),'Không cộng ảnh chụp qua ngày.'),('Trứng tháng',fmt(m['tong_trung'],'quả'),'Phát sinh được cung cấp.'),('Chất lượng',fmt(m['ty_le_dat_pct'],'%'),'Tổng trứng đạt / tổng trứng.')])
    latest=current['dan_ga'].ngay.max()
    slide=deck.slide('02 | Tổng đàn',f'Ảnh chụp ngày {latest:%d/%m/%Y}: {fmt(m["tong_dan"],"con")}.')
    cats,values=grouped_chart(summary['Tong_dan'],['phap_nhan','khu'],'so_con');deck.chart(slide,cats,[('Tổng đàn',values)],'con')
    moves=a['movements']
    slide=deck.slide('03 | Biến động đàn','Biến động chỉ được giải thích khi có luồng nhập, xuất, chết và loại.')
    if moves.empty: deck.bullets(slide,['Chưa có bảng bien_dong; không suy ra nguyên nhân từ hai ảnh chụp đàn.'])
    else:
        cols=['dau_ky','nhap','xuat','chet','loai','cuoi_ky']; values=[moves[c].sum() for c in cols]
        deck.chart(slide,['Đầu kỳ','Nhập','Xuất','Chết','Loại','Cuối kỳ'],[('Số con (xuất/chết/loại là luồng ra)',values)],'con')
    slide=deck.slide('04 | Sản lượng trứng',f'Tổng nguồn tháng: {fmt(m["tong_trung"],"quả")}; chưa mặc định đủ ngày.')
    cats,values=grouped_chart(current['trung'],['phap_nhan','khu'],'tong_trung');deck.chart(slide,cats,[('Sản lượng',values)],'quả')
    slide=deck.slide('05 | Tỷ lệ đẻ','Tỷ lệ đẻ cần mẫu số gà đẻ-ngày cùng phạm vi sản lượng.')
    deck.cards(slide,[('Tỷ lệ đẻ',fmt(m['ty_le_de_pct'],'%'),'Tổng trứng / tổng gà đẻ-ngày × 100.'),('Căn cứ mẫu số','Khớp khóa' if m['ty_le_de_pct'] is not None else 'Chưa đủ căn cứ','Ghép đúng ngày / pháp nhân / khu; không dùng tổng đàn thay gà đẻ.')])
    slide=deck.slide('06 | Chất lượng trứng',f'Tỷ lệ đạt theo nguồn cung cấp: {fmt(m["ty_le_dat_pct"],"%")}.')
    deck.chart(slide,['Trứng đạt','Trứng không đạt'],[('Chất lượng',[m['trung_dat'],m['tong_trung']-m['trung_dat']])],'quả')
    slide=deck.slide('07 | Cám sử dụng',f'Tổng lượng cám ghi nhận: {fmt(m["cam_kg"],"kg")}.')
    cats,values=grouped_chart(current['cam'],['phap_nhan','khu','loai_cam'],'khoi_luong_kg');deck.chart(slide,cats,[('Khối lượng',values)],'kg')
    slide=deck.slide('08 | KPI sản xuất','Chênh lệch và tỷ lệ hoàn thành chỉ tính khi có kế hoạch đúng kỳ, đúng đơn vị.')
    rows=[]
    for r in a['comparison'].to_dict('records'):
        rows.append([r['ten']+' ('+r['don_vi']+')',fmt(r['ke_hoach']),fmt(r['thuc_te']),fmt(r['chenh_lech'], 'đ.%' if r['don_vi']=='%' else ''),fmt(r['hoan_thanh_pct'],'%')])
    deck.table(slide,['Chỉ số / đơn vị','Kế hoạch','Thực tế','Chênh lệch','TH/KH'],rows)
    slide=deck.slide('09 | KPI nhân sự','Xếp hạng theo trọng số khai báo; chưa thay thế chính sách nhân sự được phê duyệt.')
    ranking=summary['Xep_hang_KPI']
    deck.table(slide,['Hạng','Mã nhân viên','Tên nhân viên','Điểm trọng số'],[[r.hang,r.ma_nv,r.ho_ten,fmt(r.diem_trong_so)] for r in ranking.head(8).itertuples()])
    deck.box(slide,.65,6.38,12,.25,f'Hiển thị {min(8,len(ranking))}/{len(ranking)} nhân viên; đầy đủ trong phụ lục. Đồng điểm đồng hạng.',10,MUTED)
    ytd=a['ytd']; complete=(ytd.ngay_trung==ytd.ngay_thang).all()
    slide=deck.slide(f'10 | Lũy kế {int(month[5:])} tháng',f'Trứng cộng từ nguồn: {fmt(ytd.trung_qua.sum(min_count=1),"quả")}. '+('Có dữ liệu trứng mỗi ngày; vẫn cần xác nhận phạm vi.' if complete else 'Chưa đủ phủ ngày; lũy kế chưa đầy đủ.'))
    deck.chart(slide,ytd.thang.tolist(),[('Trứng thực tế',ytd.trung_qua.tolist()),('Kế hoạch trứng',ytd.ke_hoach_trung.tolist())],'quả',XL_CHART_TYPE.LINE_MARKERS)
    deck.box(slide,.65,6.35,12,.26,f'Cám lũy kế từ nguồn: {fmt(ytd.cam_kg.sum(min_count=1),"kg")}. Chi tiết phủ ngày và kế hoạch trong phụ lục.',10,MUTED)
    # Narrative is quoted as source statements, never treated as verified causality.
    def quoted(group):
        return [f"Theo {r['nguon']}: {r['noi_dung']} [Căn cứ khai báo: {r['can_cu']}; chưa xác minh]" for r in narratives if r['nhom']==group]
    slide=deck.slide('11 | Tồn tại & căn cứ nguyên nhân','Chênh lệch là quan sát; quan hệ nhân quả cần bằng chứng riêng.')
    deck.bullets(slide,a['observations'][:2]+quoted('nguyen_nhan')+quoted('ton_tai')+a['observations'][2:])
    slide=deck.slide('12 | Giải pháp đề xuất','Ưu tiên hoàn thiện nguồn và kiểm chứng nguyên nhân trước khi giao biện pháp.')
    deck.bullets(slide,a['solutions'][:2]+quoted('giai_phap')+a['solutions'][2:])
    slide=deck.slide('13 | Kiến nghị Ban Lãnh đạo','Phê duyệt dữ liệu và chính sách KPI trước khi phát hành báo cáo chính thức.')
    deck.bullets(slide,quoted('kien_nghi')+['Đề xuất: giao đầu mối xác nhận phạm vi dữ liệu và đối soát chênh lệch; chưa xác lập người phụ trách hoặc thời hạn.',
        'Đề xuất: duyệt công thức KPI, chiều đánh giá và trọng số trước khi áp dụng cho nhân sự.',
        'Chưa đủ căn cứ đề xuất ngân sách hoặc mục tiêu mới; không tự đặt số cam kết.'])
    for slide in deck.prs.slides:
        slide.notes_slide.notes_text_frame.text = ('Nguồn: xem sheet Nguon (SHA-256) và Word_nguon trong Doi_soat_BOD.xlsx. '
            'Số liệu thực tế và kế hoạch chỉ phản ánh nguồn cung cấp. Tỷ lệ TH/KH không phải điểm hiệu suất cám. '
            'Tổng đàn dùng ảnh chụp mới nhất; trứng, cám cộng phát sinh; tỷ lệ đẻ dùng tổng gà đẻ-ngày. '
            'KPI nhân sự: chiều tăng TH/MT, chiều giảm MT/TH, nhân trọng số; không giới hạn trần. '
            'Nhận định nguyên nhân từ Word chưa xác minh; đề xuất không phải quyết định đã phê duyệt. '
            + ('DỮ LIỆU MẪU GIẢ LẬP.' if sample else 'TÀI LIỆU NỘI BỘ.'))
    deck.validate(); return deck


def audit_workbook(path,a,extra,narratives,issues,sources):
    sheets={'KPI_san_xuat':a['comparison'],'KPI_nhan_su':a['summary']['Xep_hang_KPI'],
            'Luy_ke':a['ytd'],'Phu_ngay_thang':a['coverage'],'Kiem_tra':issues,
            'Nguon':pd.DataFrame(sources),'Word_nguon':pd.DataFrame(narratives),
            'Ton_tai':pd.DataFrame({'noi_dung':a['observations']}),
            'Giai_phap':pd.DataFrame({'noi_dung':a['solutions']}),
            'Phuong_phap':a['summary']['Ghi_chu'],
            'Tong_dan':a['summary']['Tong_dan'],'Trung':a['summary']['San_luong_trung'],
            'Cam':a['summary']['Su_dung_cam'],**extra}
    with pd.ExcelWriter(path,engine='openpyxl') as writer:
        for key,df in sheets.items():
            df.to_excel(writer,sheet_name=key,index=False)
            ws=writer.sheets[key]; ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
            for row in ws:
                for cell in row:
                    if isinstance(cell.value,str): cell.data_type='s'


def run_bod(folder,output,month='2026-10',sample=False,pdf=False,recursive=False,confirmed_real=False,
            font='Montserrat',fallback_fonts=None,libreoffice_path=None):
    folder=Path(folder).expanduser().resolve();output=Path(output).expanduser().resolve()
    if folder==output or output.is_relative_to(folder) or folder.is_relative_to(output):
        raise ValueError('Thư mục đầu vào và đầu ra phải tách biệt, không lồng nhau.')
    files=discover_files(folder,recursive)
    if not files:
        raise ValueError('Không có dữ liệu thật: không tìm thấy Excel/Word. Không tự dùng dữ liệu mẫu.')
    output.mkdir(parents=True,exist_ok=True)
    name=f'Bao_cao_BOD_PCN_T{month[5:]}_{month[:4]}'
    targets=[output/(name+'.pptx'),output/'Doi_soat_BOD.xlsx',output/'Kiem_tra_BOD.json']
    if any(p.exists() for p in targets) or (output/(name+'.pdf')).exists():
        raise ValueError('Đầu ra đã tồn tại; chọn thư mục mới để không ghi đè báo cáo.')
    base,extra,narratives,issues,sources=load(folder,month,recursive,files)
    if issues.muc_do.eq('LOI').any():
        issues.to_json(output/'Loi_BOD.json',orient='records',force_ascii=False,indent=2)
        raise ValueError('Dữ liệu lỗi; xem Loi_BOD.json. Không tạo PowerPoint mới.')
    if not sample:
        protect_official_input(files,confirmed_real)
    analysis=build_analysis(base,extra,month)
    deck=create_deck(analysis,extra,narratives,month,sample,select_font(font,fallback_fonts))
    # Build all artifacts in staging before exposing a completed report.
    with tempfile.TemporaryDirectory(dir=output) as staging:
        staging=Path(staging).resolve()
        deck.prs.save(staging/targets[0].name)
        audit_workbook(staging/targets[1].name,analysis,extra,narratives,issues,sources)
        reopened=Presentation(staging/targets[0].name)
        charts=sum(s.has_chart for slide in reopened.slides for s in slide.shapes)
        checks={'month':month,'sample':sample,'slides':len(reopened.slides),'editable_charts':charts,
                'font':deck.font,'checks':deck.checks,'sources':sources,
                'limitations':['Chưa kiểm tra render trực quan trong Microsoft PowerPoint.',
                              'Nguyên nhân từ Word là nội dung nguồn chưa xác minh.',
                              'Tỷ lệ TH/KH cám không phải điểm hiệu suất; cám thấp chưa đủ chứng minh hiệu quả.',
                              'Độ phủ ngày không xác nhận đủ mọi pháp nhân/khu.'],
                'pdf':'Không yêu cầu'}
        if pdf:
            office=find_office(libreoffice_path)
            if office:
                try:
                    before_hash=hashlib.sha256((staging/targets[0].name).read_bytes()).hexdigest()
                    result=subprocess.run([office,'-env:UserInstallation='+ (staging/'lo-profile').as_uri(),'--headless','--convert-to','pdf','--outdir',str(staging),str(staging/targets[0].name)],capture_output=True,text=True,timeout=120)
                    if hashlib.sha256((staging/targets[0].name).read_bytes()).hexdigest()!=before_hash:
                        raise ValueError('Công cụ PDF đã thay đổi PowerPoint; dừng bàn giao.')
                    checks['pptx_unchanged_after_pdf']=True
                    pdf_path=staging/(name+'.pdf')
                    if result.returncode==0 and pdf_path.exists():
                        shutil.copy2(pdf_path,output/pdf_path.name);checks['pdf']='Xuất bằng LibreOffice; cần kiểm tra tương thích render.'
                    else: checks['pdf']='LibreOffice không xuất được PDF; PowerPoint vẫn giữ nguyên.'
                except (subprocess.TimeoutExpired,OSError): checks['pdf']='Xuất PDF thất bại hoặc quá thời gian; PowerPoint vẫn giữ nguyên.'
            else: checks['pdf']='Không có LibreOffice/soffice; chưa tạo PDF.'
        (staging/targets[2].name).write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
        for target in targets: shutil.copy2(staging/target.name,target)
    return checks


def main(argv=None):
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config',help='JSON cấu hình; đường dẫn tương đối tính từ thư mục cấu hình.')
    parser.add_argument('--data');parser.add_argument('--output');parser.add_argument('--month')
    for name in ['sample','pdf','recursive','confirm-real-data']:
        parser.add_argument('--'+name,action=argparse.BooleanOptionalAction,default=None)
    parser.add_argument('--font');parser.add_argument('--libreoffice-path')
    args=parser.parse_args(argv)
    try:
        config=read_config(args.config) if args.config else {}
        mapping={'data':'data_dir','output':'output_dir','confirm_real_data':'confirm_real_data'}
        for key,value in vars(args).items():
            if key!='config' and value is not None:
                config[mapping.get(key,key)]=value
        checks=run_bod(config.get('data_dir','data'),config.get('output_dir','output'),
            month=config.get('month','2026-10'),sample=config.get('sample',False),pdf=config.get('pdf',False),
            recursive=config.get('recursive',False),confirmed_real=config.get('confirm_real_data',False),
            font=config.get('font','Montserrat'),fallback_fonts=config.get('fallback_fonts'),
            libreoffice_path=config.get('libreoffice_path'))
    except (ValueError,OSError) as exc:
        print(f'Không tạo được báo cáo: {exc}');return 2
    print(f"Đã tạo {checks['slides']} slide, {checks['editable_charts']} biểu đồ chỉnh sửa được. {checks['pdf']}");return 0

if __name__=='__main__': raise SystemExit(main())
