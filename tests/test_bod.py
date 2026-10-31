import tempfile
import unittest
from pathlib import Path
import zipfile
import pandas as pd
from pptx import Presentation
from make_bod_sample import make_sample
from bod import load, build_analysis, run_bod, laying_rate, create_deck


class BODTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory();cls.root=Path(cls.temp.name)
        make_sample(cls.root/'data')
        cls.base,cls.extra,cls.narrative,cls.issues,cls.sources=load(cls.root/'data','2026-10')
        cls.a=build_analysis(cls.base,cls.extra,'2026-10')

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def test_month_ytd_and_kpi(self):
        self.assertFalse(self.issues.muc_do.eq('LOI').any())
        self.assertEqual(self.a['metrics']['tong_dan'],6000)
        self.assertEqual(self.a['metrics']['tong_trung'],31*6*800)
        self.assertEqual(self.a['metrics']['cam_kg'],31*6*120)
        self.assertEqual(self.a['metrics']['ty_le_de_pct'],80)
        self.assertEqual(self.a['metrics']['ty_le_dat_pct'],95)
        self.assertEqual(self.a['ytd'].trung_qua.sum(),304*6*800)
        comparison=self.a['comparison'].set_index('chi_so')
        self.assertAlmostEqual(comparison.loc['tong_trung','hoan_thanh_pct'],800/850*100)
        self.assertEqual(comparison.loc['cam_kg','chenh_lech'],31*6*2)
        self.assertEqual(self.a['summary']['Xep_hang_KPI'].iloc[0].ma_nv,'NV_GIA_03')
        self.assertTrue(any(n['nhom']=='nguyen_nhan' for n in self.narrative))

    def test_editable_deck_and_reopen(self):
        output=self.root/'out'
        result=run_bod(self.root/'data',output,sample=True)
        self.assertEqual(result['slides'],14);self.assertEqual(result['editable_charts'],6)
        pptx=output/'Bao_cao_BOD_PCN_T10_2026.pptx'
        prs=Presentation(pptx)
        self.assertEqual(len(prs.slides),14)
        self.assertEqual(sum(shape.has_table for slide in prs.slides for shape in slide.shapes),2)
        text=' '.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertIn('MẪU GIẢ LẬP',text)
        self.assertIn('chưa xác minh',text)
        with zipfile.ZipFile(pptx) as z:
            self.assertTrue(any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in z.namelist()))
        book=pd.read_excel(output/'Doi_soat_BOD.xlsx',sheet_name='Luy_ke')
        self.assertEqual(len(book),10)
        with self.assertRaises(ValueError):run_bod(self.root/'data',output,sample=True)

    def test_missing_denominator_and_plan(self):
        extra={k:v.copy() for k,v in self.extra.items()}
        extra['ga_de']=extra['ga_de'].iloc[:0];extra['ke_hoach']=extra['ke_hoach'].iloc[:0]
        a=build_analysis(self.base,extra,'2026-10')
        self.assertIsNone(a['metrics']['ty_le_de_pct'])
        self.assertTrue(a['comparison'].ke_hoach.isna().all())
        self.assertTrue(a['comparison'].hoan_thanh_pct.isna().all())
        self.assertEqual(len(create_deck(a,extra,[],'2026-10',True).prs.slides),14)

    def test_denominator_key_mismatch(self):
        eggs=self.a['current']['trung'];layers=self.extra['ga_de']
        self.assertIsNone(laying_rate(eggs,layers)) # extra dates must not be silently dropped
        self.assertIsNone(laying_rate(eggs,layers.iloc[:0]))

    def test_invalid_movements_units_and_duplicates(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);make_sample(folder,'2026-01')
            path=folder/'KPI_KE_HOACH_MAU_GIA_LAP.xlsx'
            sheets=pd.read_excel(path,sheet_name=None)
            sheets['ke_hoach'].loc[0,'don_vi']='tấn'
            sheets['ke_hoach']=pd.concat([sheets['ke_hoach'],sheets['ke_hoach'].iloc[[0]]])
            with pd.ExcelWriter(path) as w:
                for name,df in sheets.items():df.to_excel(w,sheet_name=name,index=False)
            production=folder/'SAN_XUAT_MAU_GIA_LAP.xlsx'
            sheets=pd.read_excel(production,sheet_name=None)
            sheets['bien_dong'].loc[0,'cuoi_ky']=999
            with pd.ExcelWriter(production) as w:
                for name,df in sheets.items():df.to_excel(w,sheet_name=name,index=False)
            _,_,_,issues,_=load(folder,'2026-01')
            errors=' '.join(issues.noi_dung)
            self.assertIn('Sai đơn vị',errors);self.assertIn('Trùng khóa',errors);self.assertIn('Không cân đối',errors)
            with self.assertRaises(ValueError):run_bod(folder,folder/'out','2026-01',True)
            self.assertFalse(list((folder/'out').glob('*.pptx')))

    def test_chart_labels_are_not_formulas(self):
        from bod import Deck
        from io import BytesIO
        from openpyxl import load_workbook
        deck=Deck('2026-10',True)
        slide=deck.slide('Kiểm thử','Nhãn nguồn phải được giữ ở dạng văn bản.')
        deck.chart(slide,['=1+1'],[('Số con',[10])],'con')
        stream=BytesIO();deck.prs.save(stream)
        with zipfile.ZipFile(stream) as package:
            embedded=next(n for n in package.namelist() if n.startswith('ppt/embeddings/'))
            wb=load_workbook(BytesIO(package.read(embedded)),data_only=False)
            self.assertNotEqual(wb.active['A2'].data_type,'f')
            self.assertEqual(wb.active['B2'].value,10)

    def test_next_month_and_partial_ytd(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);make_sample(folder,'2026-02')
            b,e,n,i,s=load(folder,'2026-02');self.assertFalse(i.muc_do.eq('LOI').any())
            a=build_analysis(b,e,'2026-02')
            self.assertEqual(a['metrics']['tong_trung'],28*6*800)
            self.assertEqual(len(a['ytd']),2)
            deck=create_deck(a,e,n,'2026-02',True)
            self.assertEqual(len(deck.prs.slides),14)
            # Removing January preserves missing values instead of inventing zeros.
            partial={k:df.loc[df.thang.eq('2026-02') if k=='kpi' else df.ngay.dt.month.eq(2)] for k,df in b.items()}
            partial_a=build_analysis(partial,e,'2026-02')
            self.assertTrue(pd.isna(partial_a['ytd'].iloc[0].trung_qua))
            self.assertEqual(partial_a['ytd'].iloc[0].ngay_trung,0)

if __name__=='__main__':unittest.main()
