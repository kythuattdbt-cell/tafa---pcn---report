"""Synthetic fixtures only: platform behavior is mocked where Windows is needed."""
import json
import os
import subprocess
import tempfile
import types
from unittest.mock import MagicMock
import unittest
from pathlib import Path
from unittest.mock import patch
import pandas as pd
from docx import Document
from bod import load, run_bod, main
from local_runtime import discover_files, protect_official_input, read_config, find_office, select_font
from make_bod_sample import make_sample


class RuntimeTests(unittest.TestCase):
    def test_config_unicode_environment_relative_paths_and_types(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);data=root/'Thư mục có dấu';data.mkdir()
            config=root/'config.local.json'
            config.write_text(json.dumps({'data_dir':'${TAFA_TEST_INPUT}', 'output_dir':'out','recursive':True,'sample':False}),encoding='utf-8-sig')
            with patch.dict(os.environ,{'TAFA_TEST_INPUT':str(data)}):
                values=read_config(config)
            self.assertEqual(values['data_dir'],str(data.resolve()))
            self.assertEqual(values['output_dir'],str(root/'out'))
            config.write_text('{"data_dir":null}')
            with self.assertRaises(ValueError):read_config(config)
            config.write_text('{"recursive":"false"}')
            with self.assertRaises(ValueError):read_config(config)
            config.write_text('{"data_dir":"${NOT_CONFIGURED_TAFA_TEST}"}')
            with self.assertRaises(ValueError):read_config(config)

    def test_recursive_sources_same_names_and_office_lock(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);make_sample(root/'Tháng 10','2026-10')
            notes=root/'Ghi chú';notes.mkdir()
            doc=Document();doc.add_paragraph('Ghi chú kiểm thử');doc.save(notes/'GHI_CHU_MAU_GIA_LAP.docx')
            (notes/'~$locked.docx').write_bytes(b'not a document')
            self.assertEqual(discover_files(root),[])
            files=discover_files(root,True);self.assertEqual(len(files),4)
            b,e,n,issues,sources=load(root,'2026-10',True)
            self.assertFalse(issues.muc_do.eq('LOI').any())
            names=[s['file'] for s in sources]
            self.assertIn('Ghi chú/GHI_CHU_MAU_GIA_LAP.docx',names)
            self.assertIn('Tháng 10/GHI_CHU_MAU_GIA_LAP.docx',names)
            self.assertTrue(b['dan_ga']['_nguon'].str.startswith('Tháng 10/').all())

    def test_official_rejects_empty_missing_and_sample(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);(root/'empty').mkdir()
            for folder in [root/'empty',root/'missing']:
                with self.assertRaises(ValueError):run_bod(folder,root/'out',confirmed_real=True)
            make_sample(root/'samples','2026-01')
            with self.assertRaisesRegex(ValueError,'giả lập'):
                run_bod(root/'samples',root/'out','2026-01',confirmed_real=True)
            self.assertFalse(list((root/'out').glob('*.pptx')))

    def test_official_detects_renamed_sample_cells_and_word(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp)
            excel=root/'input.xlsx'
            pd.DataFrame({'ten':['Nhân viên GIẢ LẬP 01']}).to_excel(excel,index=False)
            with self.assertRaisesRegex(ValueError,'giả lập'):protect_official_input([excel],True)
            word=root/'input.docx';doc=Document();doc.add_paragraph('SYNTHETIC TEST DATA');doc.save(word)
            with self.assertRaisesRegex(ValueError,'giả lập'):protect_official_input([word],True)

    def test_confirmation_required_even_without_markers(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'fixture.xlsx'
            pd.DataFrame({'fixture':[1]}).to_excel(path,index=False)
            with self.assertRaisesRegex(ValueError,'xác nhận'):protect_official_input([path])
            # Attestation is tested on a temporary synthetic fixture, never on actual personnel data.
            protect_official_input([path],True)

    def test_pdf_unavailable_and_failed_still_produce_pptx(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);make_sample(root/'data','2026-01')
            with patch('bod.find_office',return_value=None):
                checks=run_bod(root/'data',root/'no-tool','2026-01',True,True)
            self.assertIn('Không có',checks['pdf'])
            self.assertTrue((root/'no-tool'/'Bao_cao_BOD_PCN_T01_2026.pptx').exists())
            with patch('bod.find_office',return_value='soffice'), patch('bod.subprocess.run',return_value=subprocess.CompletedProcess([],1)):
                checks=run_bod(root/'data',root/'failed','2026-01',True,True)
            self.assertIn('không xuất được',checks['pdf'])
            self.assertTrue((root/'failed'/'Bao_cao_BOD_PCN_T01_2026.pptx').exists())

    def test_windows_office_programfiles_discovery_and_explicit_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            executable=Path(tmp)/'LibreOffice/program/soffice.exe'
            executable.parent.mkdir(parents=True);executable.touch()
            with patch('local_runtime.sys.platform','win32'), patch('local_runtime.shutil.which',return_value=None),patch.dict(os.environ,{'PROGRAMFILES':tmp}):
                self.assertEqual(find_office(),str(executable))
            self.assertEqual(find_office(executable),str(executable))
            self.assertIsNone(find_office(Path(tmp)/'missing.exe'))

    def test_font_detection_uses_files_not_subprocess(self):
        # Real font inspection runs without an external Linux command.
        with patch('subprocess.run',side_effect=AssertionError('No external font command allowed')):
            font=select_font('Nonexistent-TAFA-Font',['Arial','DejaVu Sans'])
        self.assertIn(font,['Arial','DejaVu Sans'])

    def test_windows_registry_font_family_selection(self):
        registry=types.SimpleNamespace(HKEY_LOCAL_MACHINE=1,HKEY_CURRENT_USER=2,
            OpenKey=MagicMock(),QueryInfoKey=lambda key:(0,1,0),
            EnumValue=lambda key,index:('Montserrat (TrueType)','Montserrat-Regular.ttf',1))
        with patch('local_runtime.sys.platform','win32'),patch.dict('sys.modules',{'winreg':registry}), \
             patch('PIL.ImageFont.truetype',return_value=types.SimpleNamespace(getname=lambda:('Montserrat','Regular'))):
            self.assertEqual(select_font('Montserrat',['Arial']),'Montserrat')
        with patch('PIL.ImageFont.truetype',return_value=types.SimpleNamespace(getname=lambda:('Arial','Regular'))):
            self.assertEqual(select_font('Unavailable',['Arial']),'Arial')

    def test_cli_config_sample_opt_in_and_override(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);make_sample(root/'data','2026-01')
            config=root/'config.local.json'
            config.write_text(json.dumps({'data_dir':'data','output_dir':'out','month':'2026-01','sample':False,'recursive':True}),encoding='utf-8')
            self.assertEqual(main(['--config',str(config)]),2)
            self.assertEqual(main(['--config',str(config),'--sample']),0)
            self.assertTrue((root/'out'/'Bao_cao_BOD_PCN_T01_2026.pptx').exists())

    def test_output_cannot_be_nested_in_recursive_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(ValueError,'không lồng nhau'):
                run_bod(tmp,Path(tmp)/'reports',recursive=True,sample=True)

if __name__=='__main__':unittest.main()
