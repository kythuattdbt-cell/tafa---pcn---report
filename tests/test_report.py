import tempfile
import unittest
from pathlib import Path
import pandas as pd
from docx import Document
from openpyxl import load_workbook
from report import read_inputs, analyze, run


def sample_frames():
    return {
        'dan_ga':pd.DataFrame([
            ['2026-10-01','PHAP_NHAN_MAU','A',100],
            ['2026-10-31','PHAP_NHAN_MAU','A',90],
            ['2026-10-31','PHAP_NHAN_MAU','B',60]],columns=['ngay','phap_nhan','khu','so_con']),
        'trung':pd.DataFrame([
            ['2026-10-01','PHAP_NHAN_MAU','A',100,90],
            ['2026-10-02','PHAP_NHAN_MAU','A',200,160],
            ['2026-10-01','PHAP_NHAN_MAU','B',0,0]],columns=['ngay','phap_nhan','khu','tong_trung','trung_dat']),
        'cam':pd.DataFrame([['2026-10-01','PHAP_NHAN_MAU','A','CAM_MAU',25.5]],columns=['ngay','phap_nhan','khu','loai_cam','khoi_luong_kg']),
        'kpi':pd.DataFrame([
            ['2026-10','NV_GIA_01','Nhan vien gia lap 01','San luong',60,100,120,'tang'],
            ['2026-10','NV_GIA_01','Nhan vien gia lap 01','Tieu hao',40,10,8,'giam'],
            ['2026-10','NV_GIA_02','Nhan vien gia lap 02','San luong',100,100,100,'tang']],columns=['thang','ma_nv','ho_ten','chi_tieu','trong_so','muc_tieu','thuc_hien','chieu'])}


def write_sample(folder,frames=None):
    folder = Path(folder); folder.mkdir(parents=True,exist_ok=True)
    with pd.ExcelWriter(folder/'MAU_GIA_LAP.xlsx') as w:
        for name,df in (frames or sample_frames()).items(): df.to_excel(w,sheet_name=name,index=False)


class ReportingTests(unittest.TestCase):
    def test_end_to_end(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); write_sample(root/'data')
            data,issues=read_inputs(root/'data','2026-10')
            self.assertTrue(issues.empty)
            result=analyze(data,'2026-10')
            self.assertEqual(result['Tong_dan'].so_con.sum(),150)
            self.assertAlmostEqual(result['San_luong_trung'].iloc[0].ty_le_dat_pct,250/300*100)
            self.assertTrue(pd.isna(result['San_luong_trung'].iloc[1].ty_le_dat_pct))
            self.assertEqual(result['Xep_hang_KPI'].iloc[0].diem_trong_so,122)
            self.assertTrue(run(root/'data',root/'out',sample=True))
            book=load_workbook(next((root/'out').glob('*.xlsx')))
            self.assertIn('Kiem_tra_du_lieu',book.sheetnames)
            doc=Document(next((root/'out').glob('*.docx')))
            self.assertTrue(any('MẪU GIẢ LẬP' in p.text for p in doc.paragraphs))
            self.assertEqual(len(doc.tables),4)

    def test_invalid_data_blocks_reports(self):
        cases=[('trung','trung_dat',101),('cam','khoi_luong_kg',-1),
               ('dan_ga','so_con',1.5),('dan_ga','khu','G'),
               ('dan_ga','ngay','2026-09-01'),('kpi','trong_so',1),
               ('kpi','muc_tieu',0),('kpi','ho_ten',None),
               ('cam','khoi_luong_kg',float('inf'))]
        for sheet,col,value in cases:
            with self.subTest(sheet=sheet,col=col,value=value), tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp); frames=sample_frames(); frames[sheet] = frames[sheet].astype(object); frames[sheet].loc[0,col]=value
                write_sample(root/'data',frames)
                self.assertFalse(run(root/'data',root/'out'))
                self.assertTrue((root/'out'/'loi_du_lieu.csv').exists())
                self.assertFalse(list((root/'out').glob('*.docx')))

    def test_multiple_workbooks_and_duplicate_detection(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp); write_sample(root)
            with pd.ExcelWriter(root/'other.xlsx') as w:
                sample_frames()['cam'].to_excel(w,sheet_name='cam',index=False)
            _,issues=read_inputs(root,'2026-10')
            self.assertTrue(issues.noi_dung.str.contains('Trùng khóa').any())

    def test_missing_groups(self):
        with tempfile.TemporaryDirectory() as tmp:
            _,issues=read_inputs(tmp,'2026-10')
            self.assertEqual(sum(issues.muc_do.eq('LOI')),5)

if __name__=='__main__': unittest.main()
