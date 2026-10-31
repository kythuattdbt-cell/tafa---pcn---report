"""Cross-platform local configuration, file discovery and input safeguards."""
import json
import os
import re
import shutil
import sys
import unicodedata
from pathlib import Path

EXTENSIONS = {'.xlsx', '.xlsm', '.xls', '.docx'}


def discover_files(folder, recursive=False):
    root = Path(folder).expanduser().resolve()
    if not root.is_dir():
        raise ValueError('Không tìm thấy thư mục đầu vào. Kiểm tra đường dẫn và đồng bộ OneDrive.')
    files = []
    def failed(error):
        raise ValueError('Không truy cập được thư mục đầu vào; kiểm tra quyền và trạng thái OneDrive.') from error
    for directory, children, names in os.walk(root, onerror=failed, followlinks=False):
        children[:] = sorted(c for c in children if c not in {'.git', '.venv', 'venv', '__pycache__'}
                             and not (Path(directory)/c).is_symlink()
                             and not getattr(Path(directory)/c, 'is_junction', lambda: False)())
        for name in sorted(names):
            path = Path(directory) / name
            if name.startswith('~$') or path.suffix.lower() not in EXTENSIONS:
                continue
            if not path.resolve().is_relative_to(root):
                raise ValueError('Tệp liên kết trỏ ra ngoài thư mục đầu vào; cần đối soát nguồn.')
            files.append(path)
        if not recursive:
            break
    return sorted(files)


def source_name(path, root):
    return Path(path).relative_to(Path(root).expanduser().resolve()).as_posix()


def normalized(text):
    value = unicodedata.normalize('NFKD', str(text))
    return ''.join(c for c in value if not unicodedata.combining(c)).upper()


def is_sample_text(text):
    value = normalized(text)
    return bool(re.search(r'GIA[ _-]*LAP|SYNTHETIC|NV_GIA_|PHAP_NHAN_MAU|CAM_MAU', value))


def protect_official_input(files, confirmed_real=False):
    """Marker detection plus explicit attestation; never certifies factual authenticity."""
    from docx import Document
    from openpyxl import load_workbook
    if not files:
        raise ValueError('Không có dữ liệu thật: không tìm thấy Excel/Word trong thư mục đầu vào.')
    for path in files:
        if is_sample_text(path.name):
            raise ValueError('Phát hiện dữ liệu giả lập. Chế độ chính thức không chấp nhận dữ liệu mẫu.')
        if path.suffix.lower() in {'.xlsx', '.xlsm'}:
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    if is_sample_text(sheet.title):
                        raise ValueError('Phát hiện nhãn giả lập trong workbook; dừng chế độ chính thức.')
                    for row in sheet.iter_rows(values_only=True):
                        if any(isinstance(v, str) and is_sample_text(v) for v in row):
                            raise ValueError('Phát hiện nội dung giả lập trong Excel; dừng chế độ chính thức.')
            finally:
                book.close()
        elif path.suffix.lower() == '.docx':
            doc = Document(path)
            texts = [p.text for p in doc.paragraphs]
            texts += [c.text for t in doc.tables for row in t.rows for c in row.cells]
            if any(is_sample_text(t) for t in texts):
                raise ValueError('Phát hiện nội dung giả lập trong Word; dừng chế độ chính thức.')
    if not confirmed_real:
        raise ValueError('Chế độ chính thức yêu cầu xác nhận nguồn thật: --confirm-real-data hoặc confirm_real_data=true trong cấu hình riêng.')


def select_font(preferred='Montserrat', fallbacks=None):
    """Inspect font files/Windows registry without invoking platform shell commands."""
    from PIL import ImageFont
    names = [preferred] + (fallbacks or ['Aptos', 'Calibri', 'Arial', 'DejaVu Sans'])
    directories = [Path.home()/'.fonts', Path.home()/'.local/share/fonts', Path('/usr/share/fonts'),
                   Path('/usr/local/share/fonts'), Path('/Library/Fonts'), Path.home()/'Library/Fonts']
    candidates = []
    if sys.platform == 'win32':
        directories += [Path(os.environ.get('WINDIR', 'C:/Windows'))/'Fonts',
                        Path(os.environ.get('LOCALAPPDATA', str(Path.home())))/'Microsoft/Windows/Fonts']
        import winreg
        for hive in [winreg.HKEY_LOCAL_MACHINE, winreg.HKEY_CURRENT_USER]:
            try:
                with winreg.OpenKey(hive, r'SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts') as key:
                    for i in range(winreg.QueryInfoKey(key)[1]):
                        _, value, _ = winreg.EnumValue(key, i)
                        if isinstance(value, str):
                            p = Path(value)
                            candidates.append(p if p.is_absolute() else directories[-2]/p)
            except OSError:
                pass
    for directory in directories:
        if directory.is_dir():
            for extension in ['*.ttf', '*.otf', '*.ttc']:
                candidates.extend(directory.rglob(extension))
    installed = set()
    for path in set(candidates):
        try:
            installed.add(ImageFont.truetype(str(path), 12).getname()[0].casefold())
        except (OSError, ValueError):
            continue
    return next((n for n in names if n.casefold() in installed), 'Arial')


def find_office(configured=None):
    if configured:
        p = Path(os.path.expandvars(str(configured))).expanduser()
        return str(p.resolve()) if p.is_file() else None
    found = shutil.which('libreoffice') or shutil.which('soffice')
    if found:
        return found
    if sys.platform == 'win32':
        for key in ['PROGRAMFILES', 'PROGRAMFILES(X86)']:
            base = os.environ.get(key)
            if base:
                path = Path(base)/'LibreOffice/program/soffice.exe'
                if path.is_file():
                    return str(path)
    return None


def read_config(path):
    path = Path(path).resolve()
    with path.open(encoding='utf-8-sig') as stream:
        config = json.load(stream)
    allowed = {'data_dir', 'output_dir', 'month', 'recursive', 'sample', 'pdf',
               'confirm_real_data', 'font', 'fallback_fonts', 'libreoffice_path'}
    if not isinstance(config, dict) or set(config)-allowed:
        raise ValueError('Cấu hình không hợp lệ hoặc có khóa không hỗ trợ.')
    for key in ['recursive', 'sample', 'pdf', 'confirm_real_data']:
        if key in config and not isinstance(config[key], bool):
            raise ValueError(f'{key} phải là true/false.')
    for key in ['data_dir', 'output_dir', 'month', 'font', 'libreoffice_path']:
        if key in config and not (key == 'libreoffice_path' and config[key] is None) and (not isinstance(config[key], str) or not config[key].strip()):
            raise ValueError(f'{key} phải là chuỗi không trống.')
    if 'fallback_fonts' in config and (not isinstance(config['fallback_fonts'], list) or not all(isinstance(s,str) and s.strip() for s in config['fallback_fonts'])):
        raise ValueError('fallback_fonts phải là danh sách tên font.')
    for key in ['data_dir', 'output_dir', 'libreoffice_path']:
        if not config.get(key):
            continue
        value = os.path.expandvars(config[key])
        if re.search(r'\$\{|\$[A-Za-z_]|%[A-Za-z_][A-Za-z0-9_]*%', value):
            raise ValueError(f'Biến môi trường của {key} chưa được thiết lập.')
        p = Path(value).expanduser()
        config[key] = str((path.parent/p).resolve() if not p.is_absolute() else p.resolve())
    return config
